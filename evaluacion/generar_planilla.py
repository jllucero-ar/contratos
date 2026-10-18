"""
generar_planilla.py — evaluacion/generar_planilla.py
Genera una planilla Excel para que un abogado revise las anotaciones de riesgos del golden set y las
citas del CCyCN que devolvió el modelo en una corrida de evaluación.

Uso:
    python evaluacion/generar_planilla.py [reporte.json]     # por defecto, el reporte más reciente
Salida: evaluacion/revision/planilla_revision_juridica.xlsx  (carpeta ignorada por git)
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).parent
REPORTES = AQUI / "reportes"
GOLDEN = AQUI / "golden"
SALIDA = AQUI / "revision" / "planilla_revision_juridica.xlsx"
NEUTROS = {729, 961}

FUENTE = "Arial"
F_NORMAL = Font(name=FUENTE, size=10)
F_NEGRITA = Font(name=FUENTE, size=10, bold=True)
F_HEADER = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
F_TITULO = Font(name=FUENTE, size=14, bold=True)
FILL_HEADER = PatternFill("solid", start_color="1F3864")
FILL_ENTRADA = PatternFill("solid", start_color="FFFF99")  # celdas que completa quien revisa
FILL_AUTO = PatternFill("solid", start_color="F2F2F2")
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
ENVOLVER = Alignment(wrap_text=True, vertical="top")

VEREDICTO_CITA = ["Correcta y pertinente", "Existe pero no es pertinente", "Incorrecta / mal aplicada", "No puedo determinar"]
SI_NO = ["Sí", "No"]
SI_NO_MATIZ = ["Sí", "No", "Sí (con matices)"]


def ultimo_reporte() -> Path:
    candidatos = sorted(REPORTES.glob("eval_*.json"))
    if not candidatos:
        raise SystemExit("No hay reportes en evaluacion/reportes/. Corré antes: python evaluacion/evaluar.py")
    return candidatos[-1]


def cargar_datos(path_reporte: Path):
    rep = json.loads(path_reporte.read_text(encoding="utf-8"))
    ref = json.loads((AQUI / "referencia_ccycn.json").read_text(encoding="utf-8"))["articulos"]
    esperados = {}
    for f in GOLDEN.glob("*.esperado.json"):
        e = json.loads(f.read_text(encoding="utf-8"))
        esperados[e["id"]] = {r["id"]: r for r in e["riesgos_esperados"]}
    return rep, ref, esperados


def _rubro(ref: dict, art: int) -> str:
    return ref.get(str(art)) or "(sin rubro verificado: consultar el código)"


def _hoja(wb, titulo, columnas, anchos, filas, entrada_desde):
    """Crea una hoja con encabezado, filas y las columnas desde `entrada_desde` (1-based) como celdas a completar."""
    ws = wb.create_sheet(titulo)
    ws.append(columnas)
    for i, c in enumerate(ws[1], start=1):
        c.font, c.fill, c.alignment, c.border = F_HEADER, FILL_HEADER, Alignment(wrap_text=True, vertical="center"), BORDE
        ws.column_dimensions[c.column_letter].width = anchos[i - 1]
    for fila in filas:
        ws.append(fila)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.font, c.alignment, c.border = F_NORMAL, ENVOLVER, BORDE
            c.fill = FILL_ENTRADA if c.column >= entrada_desde else FILL_AUTO
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 42
    return ws


def _lista(ws, valores, col_letra, n_filas):
    if n_filas < 1:
        return
    dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letra}2:{col_letra}{n_filas + 1}")


def construir(rep: dict, ref: dict, esperados: dict, nombre_reporte: str) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    # --- 1. Instrucciones
    ws = wb.create_sheet("Instrucciones")
    ws.column_dimensions["A"].width = 120
    texto = [
        ("Revisión jurídica del golden set de evaluación", F_TITULO),
        (f"Corrida evaluada: {rep['proveedor']} / {rep['modelo']} · prompt {rep['prompt_version']} · {rep['timestamp']} "
         f"(archivo {nombre_reporte})", F_NORMAL),
        ("", F_NORMAL),
        ("Qué se te pide", F_NEGRITA),
        ("Los 5 contratos del golden set son ficticios. Para cada uno, el proyecto guarda (a) los riesgos que debería marcar un revisor "
         "cuidadoso y (b) qué artículos del CCyCN serían respaldo aceptable. Esas anotaciones son un BORRADOR escrito con ayuda de IA: "
         "necesitan tu criterio. Tu revisión decide cuánto valen las métricas de riesgos y citas.", F_NORMAL),
        ("", F_NORMAL),
        ("Cómo completar", F_NEGRITA),
        ("• Las celdas AMARILLAS son para vos; las grises vienen del sistema. Varias tienen lista desplegable.", F_NORMAL),
        ("• Hoja 'Riesgos esperados': ¿es realmente un riesgo? ¿qué artículos serían el respaldo correcto? Mirá también qué dijo el modelo.", F_NORMAL),
        ("• Hoja 'Citas a revisar': para cada artículo citado por el modelo, ¿es correcto y pertinente a ese riesgo?", F_NORMAL),
        ("• Hoja 'Riesgos adicionales': riesgos que el modelo planteó y que NO estaban anotados. Algunos pueden ser válidos (y valiosos).", F_NORMAL),
        ("• Hoja 'Resumen': se calcula sola con lo que completes.", F_NORMAL),
        ("", F_NORMAL),
        ("Advertencias", F_NEGRITA),
        ("• Los títulos de artículos se tomaron del texto publicado en Infoleg (Ley 26.994) y NO se verificó si reflejan reformas posteriores "
         "(p. ej. Ley 27.551 o DNU 70/2023 en locaciones). Por favor contrastá con el texto vigente.", F_NORMAL),
        ("• Un artículo marcado 'a revisar' no está necesariamente mal: significa que no figuraba en la lista anotada.", F_NORMAL),
        ("• El modelo respondió una sola vez por contrato; otra corrida puede dar riesgos y citas distintos.", F_NORMAL),
        ("• Los art. 729 y 961 (buena fe) se toleran en cualquier riesgo y aparecen como 'neutro'.", F_NORMAL),
        ("", F_NORMAL),
        ("Fuente de consulta: https://servicios.infoleg.gob.ar/infolegInternet/anexos/235000-239999/235975/norma.htm", F_NORMAL),
    ]
    for i, (t, f) in enumerate(texto, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")

    # --- 2. Riesgos esperados
    filas = []
    for caso in rep["casos"]:
        for d in caso["riesgos"]["detalle"]:
            anot = esperados.get(caso["id"], {}).get(d["id"], {})
            arts = "; ".join(f"{a} – {_rubro(ref, a)}" for a in anot.get("articulos", [])) or "(no se exige cita)"
            modelo = "\n\n".join(d.get("textos_modelo", [])) or "— el modelo NO marcó este riesgo —"
            filas.append([caso["id"], d["id"], d["descripcion"], ", ".join(anot.get("alguno", [])), arts,
                          "Sí" if d["detectado"] else "NO", modelo, None, None, None])
    ws2 = _hoja(wb, "Riesgos esperados",
                ["Caso", "ID", "Riesgo esperado (borrador)", "Palabras clave usadas", "Artículos anotados como aceptables",
                 "¿Lo marcó el modelo?", "Qué dijo el modelo", "¿Es realmente un riesgo?",
                 "Artículos correctos (si difieren de los anotados)", "Comentario"],
                [26, 6, 40, 28, 40, 12, 70, 16, 34, 36], filas, 8)
    _lista(ws2, SI_NO_MATIZ, "H", len(filas))
    n_riesgos = len(filas)

    # --- 3. Citas a revisar
    filas = []
    for caso in rep["casos"]:
        for d in caso["riesgos"]["detalle"]:
            if not d["detectado"]:
                continue
            for art in d["citados"]:
                if art in d["respaldados"]:
                    estado = "Anotado como correcto"
                elif art in NEUTROS:
                    estado = "Neutro (tolerado)"
                else:
                    estado = "A REVISAR (no figura en la lista)"
                filas.append([caso["id"], d["id"], d["descripcion"], art, _rubro(ref, art), estado, None, None, None])
    ws3 = _hoja(wb, "Citas a revisar",
                ["Caso", "Riesgo", "Riesgo esperado", "Art. citado (CCyCN)", "Título del artículo", "Estado automático",
                 "Veredicto", "¿Agregar a la lista de aceptables?", "Comentario"],
                [26, 7, 44, 11, 40, 22, 28, 18, 36], filas, 7)
    _lista(ws3, VEREDICTO_CITA, "G", len(filas))
    _lista(ws3, SI_NO, "H", len(filas))
    n_citas = len(filas)

    # --- 4. Riesgos adicionales
    filas = []
    for caso in rep["casos"]:
        for a in caso["riesgos"]["riesgos_adicionales"]:
            filas.append([caso["id"], a["texto"], ", ".join(map(str, a["citados"])) or "—", None, None, None, None])
    ws4 = _hoja(wb, "Riesgos adicionales",
                ["Caso", "Riesgo planteado por el modelo (no anotado)", "Artículos CCyCN citados",
                 "¿Es un riesgo válido?", "¿Los artículos citados son correctos?", "¿Agregar como riesgo esperado?", "Comentario"],
                [26, 100, 18, 16, 22, 18, 36], filas, 4)
    _lista(ws4, SI_NO_MATIZ, "D", len(filas))
    _lista(ws4, ["Sí", "No", "Parcialmente", "Sin cita"], "E", len(filas))
    _lista(ws4, SI_NO, "F", len(filas))
    n_adic = len(filas)

    # --- 5. Resumen (fórmulas)
    ws5 = wb.create_sheet("Resumen")
    ws5.column_dimensions["A"].width = 58
    ws5.column_dimensions["B"].width = 14
    ws5.column_dimensions["C"].width = 14
    ws5["A1"] = "Resumen de la revisión (se calcula solo)"
    ws5["A1"].font = F_TITULO
    rc, ra, rr = f"'Citas a revisar'!$G$2:$G${max(n_citas, 2) + 1}", f"'Riesgos adicionales'!$D$2:$D${max(n_adic, 2) + 1}", \
        f"'Riesgos esperados'!$H$2:$H${max(n_riesgos, 2) + 1}"
    rev = f"'Citas a revisar'!$F$2:$F${max(n_citas, 2) + 1}"
    filas_res = [
        ("RIESGOS ESPERADOS", None, None),
        ("Total de riesgos anotados", f"=COUNTA('Riesgos esperados'!$A$2:$A${n_riesgos + 1})", None),
        ("Revisados", f"=COUNTA({rr})", "=IFERROR(B5/B4,0)"),
        ("Confirmados como riesgo (Sí / Sí (con matices))", f'=COUNTIF({rr},"Sí")+COUNTIF({rr},"Sí (con matices)")', None),
        ("Descartados (No)", f'=COUNTIF({rr},"No")', None),
        ("", None, None),
        ("CITAS DEL MODELO EN RIESGOS DETECTADOS", None, None),
        ("Total de citas a evaluar", f"=COUNTA('Citas a revisar'!$A$2:$A${n_citas + 1})", None),
        ("Revisadas", f"=COUNTA({rc})", "=IFERROR(B11/B10,0)"),
        ("Correctas y pertinentes", f'=COUNTIF({rc},"{VEREDICTO_CITA[0]}")', "=IFERROR(B12/B11,0)"),
        ("Existen pero no son pertinentes", f'=COUNTIF({rc},"{VEREDICTO_CITA[1]}")', "=IFERROR(B13/B11,0)"),
        ("Incorrectas / mal aplicadas", f'=COUNTIF({rc},"{VEREDICTO_CITA[2]}")', "=IFERROR(B14/B11,0)"),
        ("No se pudo determinar", f'=COUNTIF({rc},"{VEREDICTO_CITA[3]}")', "=IFERROR(B15/B11,0)"),
        ("De las 'A REVISAR': resultaron correctas y pertinentes",
         f'=COUNTIFS({rev},"A REVISAR*",{rc},"{VEREDICTO_CITA[0]}")', None),
        ("De las 'A REVISAR': existen pero no pertinentes o incorrectas",
         f'=COUNTIFS({rev},"A REVISAR*",{rc},"{VEREDICTO_CITA[1]}")+COUNTIFS({rev},"A REVISAR*",{rc},"{VEREDICTO_CITA[2]}")', None),
        ("", None, None),
        ("RIESGOS ADICIONALES DEL MODELO", None, None),
        ("Total", f"=COUNTA('Riesgos adicionales'!$A$2:$A${n_adic + 1})", None),
        ("Revisados", f"=COUNTA({ra})", "=IFERROR(B21/B20,0)"),
        ("Válidos (Sí / Sí (con matices))", f'=COUNTIF({ra},"Sí")+COUNTIF({ra},"Sí (con matices)")', "=IFERROR(B22/B21,0)"),
    ]
    for i, (a, b, c) in enumerate(filas_res, start=3):
        ws5.cell(row=i, column=1, value=a)
        if b is not None:
            ws5.cell(row=i, column=2, value=b)
        if c is not None:
            ws5.cell(row=i, column=3, value=c).number_format = "0.0%"
        titulo = b is None and a
        for col in (1, 2, 3):
            ws5.cell(row=i, column=col).font = F_NEGRITA if titulo else F_NORMAL
    return wb


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else ultimo_reporte()
    rep, ref, esperados = cargar_datos(path)
    wb = construir(rep, ref, esperados, path.name)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(f"✓ Planilla generada: {SALIDA}")


if __name__ == "__main__":
    main()
